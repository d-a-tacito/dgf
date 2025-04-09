import numpy as np
import matplotlib.pyplot as plt
from mpl_toolkits.mplot3d import Axes3D
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from matplotlib.patches import Ellipse

def F1(x):
    return 9*x[0]**2 + 16*x[1]**2 - 90*x[0] - 128*x[1]

def F2(x):
    return x[0]**2 + 2*x[0]*x[1] + 2*x[1]**2 + x[2]**2 - x[1]*x[2] + x[0] + 3*x[1] - x[2]

def gram_schmidt(directions):
    n = len(directions)
    ortho_directions = np.zeros_like(directions)
    
    for i in range(n):
        v = directions[i].copy()
        for j in range(i):
            v -= np.dot(directions[i], ortho_directions[j]) * ortho_directions[j]
        ortho_directions[i] = v / np.linalg.norm(v)
    
    return ortho_directions

def rosenbrock_direct_search(f, x0, epsilon=1e-6, alpha=1.3, beta=-0.5, max_iter=1000):
    n = len(x0)
    xk = np.array(x0, dtype=float)
    y1 = xk.copy()
    deltas = np.ones(n) * 0.1  # начальные длины шагов
    directions = np.eye(n)     # начальные направления - координатные оси
    
    history = []
    all_points = [y1.copy()]   # сохраняем все точки для траектории
    success_steps = []
    fail_steps = []
    
    k = 1
    while k <= max_iter:
        j = 0
        yn1 = y1.copy()
        
        while j < n:
            dj = directions[j]
            yj = yn1.copy()
            f_yj = f(yj)
            
            # Пробуем шаг
            y_new = yj + deltas[j] * dj
            f_new = f(y_new)
            
            if f_new < f_yj:  # успешный шаг
                yn1 = y_new
                deltas[j] *= alpha
                success_steps.append((yj.copy(), y_new.copy()))
                history.append({
                    'k': k,
                    'xk': xk.copy(),
                    'f_xk': f(xk),
                    'j': j+1,
                    'yj': yj.copy(),
                    'f_yj': f_yj,
                    'delta_j': deltas[j] / alpha,
                    'dj': dj.copy(),
                    'y_new': y_new.copy(),
                    'f_new': f_new,
                    'result': 'success'
                })
            else:  # неудачный шаг
                deltas[j] *= beta
                fail_steps.append((yj.copy(), y_new.copy()))
                history.append({
                    'k': k,
                    'xk': xk.copy(),
                    'f_xk': f(xk),
                    'j': j+1,
                    'yj': yj.copy(),
                    'f_yj': f_yj,
                    'delta_j': deltas[j] / beta,
                    'dj': dj.copy(),
                    'y_new': y_new.copy(),
                    'f_new': f_new,
                    'result': 'fail'
                })
            
            all_points.append(y_new.copy())
            j += 1
        
        # Шаг 2 алгоритма
        if f(yn1) < f(y1):  # был хотя бы один успешный шаг
            y1 = yn1.copy()
            continue
        else:  # все шаги неудачные
            if f(yn1) < f(xk):  # был успешный шаг в течение итерации
                # Шаг 3 - формируем новые направления
                delta_k1 = yn1 - xk
                if np.linalg.norm(delta_k1) < epsilon:
                    break
                
                # Находим коэффициенты разложения
                A = directions.T
                b = delta_k1
                lambdas = np.linalg.solve(A, b)
                
                # Формируем новые направления
                new_directions = np.zeros_like(directions)
                new_directions[0] = delta_k1 / np.linalg.norm(delta_k1)
                
                for i in range(1, n):
                    new_directions[i] = directions[i]
                
                # Ортогонализация Грама-Шмидта
                directions = gram_schmidt(new_directions)
                
                # Сброс параметров
                xk = yn1.copy()
                y1 = xk.copy()
                deltas = np.ones(n) * 0.1
                k += 1
            else:  # ни одного успешного шага
                if all(np.abs(deltas) < epsilon):
                    break
                else:
                    y1 = yn1.copy()
    
    return xk, history, success_steps, fail_steps, np.array(all_points)

def plot_2d_results(success_steps, fail_steps, result_point, all_points, func_name):
    plt.figure(figsize=(12, 10))
    
    # Создаем сетку для линий уровня
    x = np.linspace(min(all_points[:, 0])-1, max(all_points[:, 0])+1, 400)
    y = np.linspace(min(all_points[:, 1])-1, max(all_points[:, 1])+1, 400)
    X, Y = np.meshgrid(x, y)
    Z = 9*X**2 + 16*Y**2 - 90*X - 128*Y
    
    # Рисуем линии уровня
    levels = np.linspace(Z.min(), Z.max(), 30)
    contour = plt.contour(X, Y, Z, levels=levels, cmap='viridis', alpha=0.5)
    plt.clabel(contour, inline=True, fontsize=8)
    
    # Рисуем эллипс вокруг минимума
    ellipse = Ellipse(xy=(5, 4), width=2, height=1.5, angle=0, 
                     edgecolor='blue', fc='None', lw=2, linestyle='--', alpha=0.7)
    plt.gca().add_patch(ellipse)
    plt.text(5, 4.5, "Минимум функции", ha='center', color='blue')
    
    # Успешные и неудачные шаги
    if success_steps:
        success_starts = np.array([s[0] for s in success_steps])
        success_ends = np.array([s[1] for s in success_steps])
        plt.quiver(success_starts[:, 0], success_starts[:, 1], 
                  success_ends[:, 0]-success_starts[:, 0], 
                  success_ends[:, 1]-success_starts[:, 1], 
                  angles='xy', scale_units='xy', scale=1, color='green', width=0.005,
                  label='Успешные шаги')
    
    if fail_steps:
        fail_starts = np.array([f[0] for f in fail_steps])
        fail_ends = np.array([f[1] for f in fail_steps])
        plt.quiver(fail_starts[:, 0], fail_starts[:, 1], 
                  fail_ends[:, 0]-fail_starts[:, 0], 
                  fail_ends[:, 1]-fail_starts[:, 1], 
                  angles='xy', scale_units='xy', scale=1, color='red', width=0.005,
                  label='Неудачные шаги')
    
    # Итоговая точка и минимум
    plt.scatter(result_point[0], result_point[1], s=200, c='magenta', marker='*', 
               edgecolors='black', linewidths=1, label='Итоговая точка')
    plt.scatter(5, 4, s=100, c='blue', marker='x', linewidths=2, label='Истинный минимум')
    
    plt.xlabel('X1')
    plt.ylabel('X2')
    plt.title(f'Метод Розенброка для функции {func_name}', pad=20)
    plt.grid(True, linestyle='--', alpha=0.7)
    plt.legend(loc='upper right', bbox_to_anchor=(1.3, 1))
    plt.tight_layout()
    plt.show()

def plot_3d_results(all_points, result_point, func_name):
    fig = plt.figure(figsize=(12, 10))
    ax = fig.add_subplot(111, projection='3d')
    
    # Траектория
    ax.plot(all_points[:, 0], all_points[:, 1], all_points[:, 2], 'o-', 
            color='blue', markersize=3, alpha=0.5, label='Траектория')
    
    # Начальная точка
    ax.scatter(all_points[0, 0], all_points[0, 1], all_points[0, 2], 
               s=100, c='green', marker='o', label='Начальная точка')
    
    # Итоговая точка
    ax.scatter(result_point[0], result_point[1], result_point[2], 
               s=200, c='red', marker='*', label='Итоговая точка')
    
    # Настройки графика
    ax.set_xlabel('X1')
    ax.set_ylabel('X2')
    ax.set_zlabel('X3')
    ax.set_title(f'Траектория метода Розенброка для функции {func_name}', pad=20)
    ax.legend()
    plt.tight_layout()
    plt.show()

def save_to_excel(history, filename):
    wb = Workbook()
    ws = wb.active
    
    headers = ['K', 'X_k', 'F(X_k)', 'j', 'y_j', 'F(y_j)', 'Δ_j', 'd_j', 'y_i + Δ_j*d_i', 'F(y_i + Δ_j*d_i)', 'Результат']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    
    for row, record in enumerate(history, 2):
        ws.cell(row=row, column=1, value=record['k'])
        ws.cell(row=row, column=2, value=str(np.round(record['xk'], 4)))
        ws.cell(row=row, column=3, value=np.round(record['f_xk'], 4))
        ws.cell(row=row, column=4, value=record['j'])
        ws.cell(row=row, column=5, value=str(np.round(record['yj'], 4)))
        ws.cell(row=row, column=6, value=np.round(record['f_yj'], 4))
        ws.cell(row=row, column=7, value=np.round(record['delta_j'], 4))
        ws.cell(row=row, column=8, value=str(np.round(record['dj'], 4)))
        ws.cell(row=row, column=9, value=str(np.round(record['y_new'], 4)))
        ws.cell(row=row, column=10, value=np.round(record['f_new'], 4))
        ws.cell(row=row, column=11, value='Успех' if record['result'] == 'success' else 'Неудача')
    
    for col in range(1, len(headers)+1):
        max_length = 0
        column = get_column_letter(col)
        for cell in ws[column]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(filename)

def main():
    print("Выберите функцию для оптимизации:")
    print("1. F1(x) = 9x1² + 16x2² - 90x1 - 128x2 (2D)")
    print("2. F2(x) = x1² + 2x1x2 + 2x2² + x3² - x2x3 + x1 + 3x2 - x3 (3D)")
    
    choice = int(input("Введите номер функции (1 или 2): "))
    epsilon = float(input("Введите значение ε (например, 0.001): "))
    x0 = input("Введите начальную точку через запятую (например, 0,3 для F1 или 0,0,0 для F2): ")
    x0 = [float(x.strip()) for x in x0.split(',')]
    
    if choice == 1:
        func = F1
        func_name = "F1"
        if len(x0) != 2:
            print("Для F1 требуется 2 координаты!")
            return
    elif choice == 2:
        func = F2
        func_name = "F2"
        if len(x0) != 3:
            print("Для F2 требуется 3 координаты!")
            return
    else:
        print("Неверный выбор функции!")
        return
    
    result, history, success_steps, fail_steps, all_points = rosenbrock_direct_search(func, x0, epsilon)
    
    print("\nРезультат оптимизации:")
    print(f"Оптимальная точка: {np.round(result, 4)}")
    print(f"Значение функции: {np.round(func(result), 4)}")
    
    # Сохранение в Excel
    excel_filename = f"rosenbrock_results_{func_name}.xlsx"
    save_to_excel(history, excel_filename)
    print(f"\nРезультаты сохранены в файл {excel_filename}")
    
    # Визуализация
    if choice == 1:
        plot_2d_results(success_steps, fail_steps, result, all_points, func_name)
    else:
        plot_3d_results(all_points, result, func_name)

if __name__ == "__main__":
    main()